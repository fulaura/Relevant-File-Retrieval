import docx
from pathlib import Path
import os
from datetime import datetime

def parse_docx(file_path):
    doc = docx.Document(file_path)
    
    # Extract text
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    # Extract tables
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)

    # File metadata
    file_name = Path(file_path).stem
    file_type = "docx"
    stat = os.stat(file_path)
    created_date = datetime.fromtimestamp(stat.st_ctime)
    modified_date = datetime.fromtimestamp(stat.st_mtime)

    return {
        "file_name": file_name,
        "file_type": file_type,
        "created_date": created_date,
        "modified_date": modified_date,
        "content": text.strip(),
        "metadata": None,
        "path": str(Path(file_path).resolve()),
        "embedding": None
    }
